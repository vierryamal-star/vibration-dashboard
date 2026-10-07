import streamlit as st
import sys
import os
import re
import html
import datetime
from bisect import bisect_left, bisect_right

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
from utils import render_page_header, render_app_sidebar, GLOBAL_UI_CSS

try:
    st.set_page_config(
        page_title="Datasheet Pompa — PLTU TBK",
        page_icon="📋",
        layout="wide",
        initial_sidebar_state="expanded",
    )
except Exception:
    pass

st.markdown(
    "<style>[data-testid='stSidebarNav'] { display: none !important; } section[data-testid='stSidebar'] > div:first-child { padding-top: 1rem; }</style>",
    unsafe_allow_html=True,
)
st.markdown(GLOBAL_UI_CSS, unsafe_allow_html=True)

render_app_sidebar()
render_page_header("📋 Datasheet Pompa & Peralatan Utama PLTU TBK")
st.caption("Spesifikasi teknis peralatan langsung dari file datasheet PLTU Tanjung Balai Karimun — satu tab per sheet Excel.")

# Sheet yang tidak ingin ditampilkan (cocokkan sebagian nama, huruf kecil).
# Sheet yang di-hidden di Excel (mis. Sheet2) otomatis dilewati.
EXCLUDED_SHEETS = ["duplex"]

# ── Cari lokasi file Excel ──────────────────────────────────────────────────
current_dir = os.path.dirname(os.path.abspath(__file__))
root_dir = os.path.dirname(current_dir)

PREFERRED_NAMES = [
    "Datasheet Pompa PLTU TBK (1).xlsx",
    "Datasheet Pompa PLTU TBK.xlsx",
]


def find_excel_file():
    for base in (root_dir, current_dir):
        for name in PREFERRED_NAMES:
            p = os.path.join(base, name)
            if os.path.exists(p):
                return p
    for base in (root_dir, current_dir):
        if not os.path.isdir(base):
            continue
        names = sorted(
            f for f in os.listdir(base)
            if f.lower().endswith((".xlsx", ".xlsm"))
            and "datasheet" in f.lower()
            and not f.startswith("~$")
        )
        if names:
            return os.path.join(base, names[0])
    return None


# ── Parser sheet Excel → struktur tabel ─────────────────────────────────────
def _fmt(v, number_format=""):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        nf = (number_format or "").lower()
        if "m" in nf and "d" not in nf:
            return v.strftime("%b %Y")
        return v.strftime("%d-%m-%Y")
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        return repr(round(v, 6))
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def _has_fill(cell):
    try:
        f = cell.fill
        if f is None or f.fill_type != "solid":
            return False
        c = f.fgColor
        if c is not None and c.type == "rgb" and c.rgb in ("00000000", "FFFFFFFF"):
            return False
        return True
    except Exception:
        return False


def _layout(raw_rows, rows, ncols, cidx):
    """Ubah item mentah menjadi baris HTML (dengan colspan/rowspan + sel kosong pengisi)."""
    out = []
    carry = {}
    for r in rows:
        occ = {c for c, n in carry.items() if n > 0}
        items = sorted(raw_rows[r], key=lambda i: i["c0"])
        cells = []
        col = 0

        def add_gap(a, b):
            run = 0
            for c in range(a, b):
                if c in occ:
                    if run:
                        cells.append((None, run, 1, "", ""))
                        run = 0
                else:
                    run += 1
            if run:
                cells.append((None, run, 1, "", ""))

        new_carry = []
        for n_i, it in enumerate(items):
            s, e = cidx[it["c0"]], cidx[it["c1"]]
            # Teks panjang pada sel tidak di-merge "tumpah" ke kolom kosong di kanannya (seperti di Excel)
            if it["c0"] == it["c1"] and it["r2"] == r and not it["numeric"]:
                limit = cidx[items[n_i + 1]["c0"]] if n_i + 1 < len(items) else ncols
                while e + 1 < limit and (e + 1) not in occ:
                    e += 1
            if s > col:
                add_gap(col, s)
            rs = bisect_right(rows, it["r2"]) - bisect_left(rows, r)
            rs = max(rs, 1)
            cls = ("b " if it["bold"] else "") + ("f" if it["fill"] else "")
            align = it["align"] or ""
            cells.append((it["text"], e - s + 1, rs, cls.strip(), align))
            if rs > 1:
                new_carry.append((s, e, rs - 1))
            col = e + 1
        if col < ncols:
            add_gap(col, ncols)
        for c in list(carry):
            carry[c] -= 1
        for s, e, n in new_carry:
            for c in range(s, e + 1):
                carry[c] = n
        out.append(cells)
    return out


def _parse_sheet(ws):
    hidden = set()
    widths = {}
    for dim in ws.column_dimensions.values():
        if dim.min is None or dim.max is None:
            continue
        for c in range(dim.min, dim.max + 1):
            if dim.hidden:
                hidden.add(c)
            if dim.width:
                widths[c] = dim.width

    master, covered = {}, set()
    for m in ws.merged_cells.ranges:
        master[(m.min_row, m.min_col)] = (m.max_row, m.max_col)
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                if (r, c) != (m.min_row, m.min_col):
                    covered.add((r, c))

    raw_rows = {}
    for r in range(1, ws.max_row + 1):
        items = []
        for c in range(ws.min_column, ws.max_column + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(r, c)
            text = _fmt(cell.value, cell.number_format)
            if not text:
                continue
            r2, c2 = master.get((r, c), (r, c))
            vis = [x for x in range(c, c2 + 1) if x not in hidden]
            if not vis:
                continue
            items.append(dict(
                r2=r2, c0=vis[0], c1=vis[-1], text=text,
                bold=bool(cell.font and cell.font.b),
                fill=_has_fill(cell),
                align=(cell.alignment.horizontal if cell.alignment else None),
                numeric=isinstance(cell.value, (int, float, datetime.date)),
            ))
        if items:
            raw_rows[r] = items

    # Baris "EQUIPMENT NAME" → judul; baris di atasnya → info formulir
    eq_row = 0
    for r in sorted(raw_rows):
        if any(i["text"].lower().startswith("equipment name") for i in raw_rows[r]):
            eq_row = r
            break

    def value_after(label):
        for r in sorted(raw_rows):
            if r > eq_row:
                break
            its = sorted(raw_rows[r], key=lambda i: i["c0"])
            for k, it in enumerate(its):
                if it["text"].lower().rstrip(" :") == label and k + 1 < len(its):
                    return its[k + 1]["text"].lstrip(": ").strip()
        return ""

    name = ""
    for it in sorted(raw_rows.get(eq_row, []), key=lambda i: i["c0"]):
        if it["text"].startswith(":"):
            name = it["text"].lstrip(": ").strip()
            break
    info = {
        "Unit": value_after("unit"),
        "No. Formulir": value_after("no. formulir"),
        "No. Revisi": value_after("no. revisi"),
        "Tanggal": value_after("tanggal"),
        "Halaman": value_after("halaman"),
    }
    info = {k: v for k, v in info.items() if v}

    body_rows = [r for r in sorted(raw_rows) if r > eq_row]
    pic_rows = [r for r in body_rows
                if all(i["text"].strip().lower() == "picture" for i in raw_rows[r])]
    body_rows = [r for r in body_rows if r not in pic_rows]

    used = [(i["c0"], i["c1"]) for r in body_rows for i in raw_rows[r]]
    if not used:
        return dict(name=name, info=info, segments=[], images=[], hits_text=[])
    cmin = min(u[0] for u in used)
    cmax = max(u[1] for u in used)
    vis_cols = [c for c in range(cmin, cmax + 1) if c not in hidden]
    cidx = {c: i for i, c in enumerate(vis_cols)}
    ws_w = [widths.get(c, 8.43) for c in vis_cols]
    total = sum(ws_w) or 1
    col_pct = [round(w / total * 100, 2) for w in ws_w]

    # Pisahkan tabel di baris "PICTURE" supaya foto muncul di posisi yang sama
    if pic_rows:
        split = pic_rows[0]
        before = [r for r in body_rows if r < split]
        after = [r for r in body_rows if r > split]
    else:
        before, after = body_rows, []

    segments = []
    for part in (before, "IMAGES", after):
        if part == "IMAGES":
            segments.append(("images", None))
        elif part:
            segments.append(("table", _layout(raw_rows, part, len(vis_cols), cidx)))
    if not pic_rows:  # tidak ada baris PICTURE → foto di bagian bawah
        segments = [s for s in segments if s[0] != "images"] + [("images", None)]

    # Gambar (lewati logo di bagian kepala sheet)
    imgs = []
    seen = set()
    for im in getattr(ws, "_images", []):
        try:
            row0, col0 = im.anchor._from.row, im.anchor._from.col
            if row0 <= 3:
                continue
            data = im._data()
            if hash(data) in seen:
                continue
            seen.add(hash(data))
            imgs.append((row0, col0, data))
        except Exception:
            continue
    imgs.sort(key=lambda x: (x[0], x[1]))

    return dict(
        name=name, info=info, segments=segments, col_pct=col_pct,
        images=[d for _, _, d in imgs],
    )


@st.cache_data(show_spinner="Membaca datasheet Excel…")
def load_datasheet(path, mtime):
    wb = openpyxl.load_workbook(path, data_only=True)
    result = {}
    for ws in wb.worksheets:
        title = ws.title.strip()
        if ws.sheet_state != "visible":
            continue
        if any(x in title.lower() for x in EXCLUDED_SHEETS):
            continue
        result[title] = _parse_sheet(ws)
    return result


# ── Render HTML ─────────────────────────────────────────────────────────────
TABLE_CSS = """
<style>
.ds-wrap{overflow-x:auto;margin:.25rem 0 1rem 0;}
table.ds{border-collapse:collapse;table-layout:fixed;width:100%;min-width:760px;font-size:.88rem;line-height:1.35;}
table.ds td{padding:4px 7px;vertical-align:middle;overflow-wrap:anywhere;}
table.ds td.c{border:1px solid rgba(128,128,128,.35);}
table.ds td.b{font-weight:600;}
table.ds td.f{background:rgba(128,128,128,.16);}
table.ds mark{background:#ffd54f;color:#000;padding:0 2px;border-radius:2px;}
.ds-info{display:flex;flex-wrap:wrap;gap:.5rem;margin:.25rem 0 .75rem 0;}
.ds-chip{border:1px solid rgba(128,128,128,.35);border-radius:999px;padding:2px 12px;font-size:.82rem;}
</style>
"""
st.markdown(TABLE_CSS, unsafe_allow_html=True)


def _cell_html(text, kw_re):
    if kw_re is None:
        return html.escape(text).replace("$", "&#36;")
    parts = kw_re.split(text)
    out = []
    for i, p in enumerate(parts):
        e = html.escape(p).replace("$", "&#36;")
        out.append(f"<mark>{e}</mark>" if i % 2 == 1 else e)
    return "".join(out)


def table_html(rows, col_pct, kw_re):
    cols = "".join(f'<col style="width:{p}%">' for p in col_pct)
    body = []
    for cells in rows:
        tds = []
        for text, cs, rs, cls, align in cells:
            attrs = ""
            if cs > 1:
                attrs += f' colspan="{cs}"'
            if rs > 1:
                attrs += f' rowspan="{rs}"'
            if text is None:
                tds.append(f"<td{attrs}></td>")
                continue
            style = f' style="text-align:{align}"' if align in ("center", "right") else ""
            tds.append(f'<td class="c {cls}"{attrs}{style}>{_cell_html(text, kw_re)}</td>')
        body.append("<tr>" + "".join(tds) + "</tr>")
    return f'<div class="ds-wrap"><table class="ds"><colgroup>{cols}</colgroup>{"".join(body)}</table></div>'


def count_hits(sheet, kw_re):
    if kw_re is None:
        return 0
    n = 0
    for kind, rows in sheet["segments"]:
        if kind != "table":
            continue
        for cells in rows:
            for text, *_ in cells:
                if text and kw_re.search(text):
                    n += 1
    n += sum(1 for v in [sheet["name"], *sheet["info"].values()] if kw_re.search(v))
    return n


def render_images(images):
    if not images:
        return
    st.markdown("**📷 Foto / Gambar**")
    ncol = 2
    for start in range(0, len(images), ncol):
        cols = st.columns(ncol)
        for col, data in zip(cols, images[start:start + ncol]):
            with col:
                try:
                    st.image(data, use_container_width=True)
                except TypeError:
                    st.image(data, use_column_width=True)


# ── Muat data ───────────────────────────────────────────────────────────────
target_file = find_excel_file()
if not target_file:
    st.error("⚠️ File Excel datasheet (mis. `Datasheet Pompa PLTU TBK (1).xlsx`) tidak ditemukan di repositori.")
    st.info("Upload file Excel ke folder utama project (sejajar dengan folder `pages`).")
    st.stop()

sheets = load_datasheet(target_file, os.path.getmtime(target_file))
sheet_names = list(sheets.keys())

# ── Pencarian ───────────────────────────────────────────────────────────────
c_search, c_info = st.columns([3, 1.4])
with c_search:
    keyword = st.text_input(
        "🔍 Cari spesifikasi / part number / merek:",
        placeholder="Contoh: 6305, Head, Torishima, Teco...",
    )
kw = keyword.strip()
kw_re = re.compile(f"({re.escape(kw)})", re.IGNORECASE) if kw else None

hits = {n: count_hits(sheets[n], kw_re) for n in sheet_names}
with c_info:
    st.metric("Jumlah sheet", len(sheet_names))

if kw:
    found = [f"**{n}** ({h})" for n, h in hits.items() if h]
    if found:
        st.success("Ditemukan di: " + " · ".join(found))
    else:
        st.info("Tidak ada data yang cocok dengan kata kunci pencarian.")

# ── Satu tab per nama sheet ─────────────────────────────────────────────────
labels = [f"{n} 🔎{hits[n]}" if kw and hits[n] else n for n in sheet_names]
tabs = st.tabs(labels)

for tab, s_name in zip(tabs, sheet_names):
    sheet = sheets[s_name]
    with tab:
        st.subheader(sheet["name"] or s_name)
        if sheet["info"]:
            chips = "".join(
                f'<span class="ds-chip"><b>{html.escape(k)}</b>: {_cell_html(v, kw_re)}</span>'
                for k, v in sheet["info"].items()
            )
            st.markdown(f'<div class="ds-info">{chips}</div>', unsafe_allow_html=True)

        if not sheet["segments"]:
            st.info("Sheet ini tidak berisi data.")
            continue

        for kind, rows in sheet["segments"]:
            if kind == "table":
                st.markdown(table_html(rows, sheet["col_pct"], kw_re), unsafe_allow_html=True)
            else:
                render_images(sheet["images"])
